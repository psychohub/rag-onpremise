"""
RAG On-Premise — Servicio de Ingesta de Documentos
===================================================
Monitorea una carpeta y procesa automáticamente PDFs, Word y Excel.
Genera embeddings con nomic-embed-text via Ollama y los almacena en Qdrant.

Uso:
    python ingestor.py

Configuración:
    Editar config.json con las rutas y URLs de su entorno.
"""

import os
import json
import time
import logging
from pathlib import Path
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler
import pdfplumber
from docx import Document
import openpyxl
from qdrant_client import QdrantClient
from qdrant_client.models import VectorParams, Distance, PointStruct
from langchain_text_splitters import RecursiveCharacterTextSplitter
import ollama
import uuid

# ── Logging ────────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler("ingestor.log", encoding="utf-8"),
        logging.StreamHandler()
    ]
)
log = logging.getLogger(__name__)

# ── Configuración ──────────────────────────────────────────────────────────────
CONFIG_PATH = Path(__file__).parent / "config.json"

with open(CONFIG_PATH, "r", encoding="utf-8") as f:
    cfg = json.load(f)

# ── Clientes ───────────────────────────────────────────────────────────────────
qdrant = QdrantClient(url=cfg["qdrant_url"])
ollama_client = ollama.Client(host=cfg["ollama_url"])

splitter = RecursiveCharacterTextSplitter(
    chunk_size=cfg.get("chunk_size", 500),
    chunk_overlap=cfg.get("chunk_overlap", 50)
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}


# ── Funciones principales ──────────────────────────────────────────────────────

def ensure_collection() -> None:
    """Crea la colección en Qdrant si no existe."""
    existing = [c.name for c in qdrant.get_collections().collections]
    collection_name = cfg["collection_name"]

    if collection_name not in existing:
        qdrant.create_collection(
            collection_name=collection_name,
            vectors_config=VectorParams(size=768, distance=Distance.COSINE)
        )
        log.info(f"Coleccion '{collection_name}' creada.")
    else:
        log.info(f"Coleccion '{collection_name}' ya existe.")


def extract_text(filepath: str) -> str:
    """
    Extrae texto de PDF, Word o Excel.
    Retorna string vacío si no puede extraer.
    """
    ext = Path(filepath).suffix.lower()

    try:
        if ext == ".pdf":
            with pdfplumber.open(filepath) as pdf:
                pages = [p.extract_text() or "" for p in pdf.pages]
                return "\n".join(pages)

        elif ext == ".docx":
            doc = Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    line = " | ".join(
                        str(cell) for cell in row if cell is not None
                    )
                    if line.strip():
                        lines.append(line)
            return "\n".join(lines)

    except Exception as e:
        log.error(f"Error extrayendo texto de {filepath}: {e}")

    return ""


def get_embedding(text: str) -> list[float]:
    """Genera embedding usando nomic-embed-text via Ollama."""
    response = ollama_client.embeddings(
        model=cfg.get("embedding_model", "nomic-embed-text"),
        prompt=text
    )
    return response["embedding"]


def ingest_file(filepath: str) -> None:
    """
    Procesa un archivo completo:
    1. Extrae texto
    2. Divide en chunks
    3. Genera embeddings
    4. Almacena en Qdrant
    """
    log.info(f"Procesando: {filepath}")

    text = extract_text(filepath)
    if not text.strip():
        log.warning(f"Sin texto extraido: {filepath}")
        return

    chunks = splitter.split_text(text)
    filename = Path(filepath).name
    collection = cfg["collection_name"]
    points = []

    for i, chunk in enumerate(chunks):
        vector = get_embedding(chunk)
        points.append(PointStruct(
            id=str(uuid.uuid4()),
            vector=vector,
            payload={
                "filename": filename,
                "filepath": filepath,
                "chunk_index": i,
                "total_chunks": len(chunks),
                "text": chunk
            }
        ))

    qdrant.upsert(collection_name=collection, points=points)
    log.info(f"[OK] {filename}: {len(chunks)} chunks indexados.")


def initial_scan(folder: str) -> None:
    """Escanea la carpeta completa al iniciar el servicio."""
    log.info(f"Escaneo inicial: {folder}")
    count = 0

    for root, _, files in os.walk(folder):
        for file in files:
            if Path(file).suffix.lower() in SUPPORTED_EXTENSIONS:
                ingest_file(os.path.join(root, file))
                count += 1

    log.info(f"Escaneo inicial completado. {count} archivo(s) procesados.")


# ── Watchdog ───────────────────────────────────────────────────────────────────

class DocumentHandler(FileSystemEventHandler):
    """Detecta archivos nuevos o modificados y los ingesta automáticamente."""

    def on_created(self, event):
        if not event.is_directory:
            self._process(event.src_path)

    def on_modified(self, event):
        if not event.is_directory:
            self._process(event.src_path)

    def _process(self, filepath: str):
        if Path(filepath).suffix.lower() in SUPPORTED_EXTENSIONS:
            time.sleep(2)  # Esperar que termine de escribirse el archivo
            ingest_file(filepath)


# ── Entry point ────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    log.info("=== RAG Ingestor iniciando ===")
    log.info(f"Carpeta monitoreada: {cfg['documents_folder']}")
    log.info(f"Coleccion Qdrant: {cfg['collection_name']}")
    log.info(f"Modelo embeddings: {cfg.get('embedding_model', 'nomic-embed-text')}")

    # Inicializar
    ensure_collection()
    initial_scan(cfg["documents_folder"])

    # Monitorear carpeta
    observer = Observer()
    observer.schedule(
        DocumentHandler(),
        cfg["documents_folder"],
        recursive=True
    )
    observer.start()
    log.info("Monitoreando cambios en la carpeta...")

    try:
        while True:
            time.sleep(10)
    except KeyboardInterrupt:
        log.info("Deteniendo ingestor...")
        observer.stop()

    observer.join()
    log.info("Ingestor detenido.")
